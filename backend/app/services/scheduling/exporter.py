"""Schedule export service – PDF, Excel (XLSX), iCalendar."""

from __future__ import annotations

import io
from datetime import datetime

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession


async def _load_schedule(project_id: str, db: AsyncSession):
    from app.db.models import Project, Scene, ShootingDay, shooting_day_scenes
    from sqlalchemy.orm import selectinload

    proj_result = await db.execute(
        select(Project)
        .where(Project.id == project_id)
        .options(selectinload(Project.shooting_days))
    )
    project = proj_result.scalar_one_or_none()
    if not project:
        return None, []

    days_result = await db.execute(
        select(ShootingDay)
        .where(ShootingDay.project_id == project_id)
        .order_by(ShootingDay.date)
        .options(selectinload(ShootingDay.scenes))
    )
    days = list(days_result.scalars().all())
    return project, days


# ── PDF ───────────────────────────────────────────────────────────────────────

async def export_schedule_pdf(project_id: str, db: AsyncSession) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet

    project, days = await _load_schedule(project_id, db)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4)
    styles = getSampleStyleSheet()
    elements = []

    elements.append(Paragraph(f"Shooting Schedule – {project.name if project else project_id}", styles["Title"]))
    elements.append(Spacer(1, 12))

    for day in days:
        elements.append(Paragraph(f"Day {day.date.strftime('%A, %d %b %Y')}", styles["Heading2"]))
        table_data = [["Scene #", "Heading", "I/E", "D/N", "Est. (min)"]]
        for scene in day.scenes:
            table_data.append([
                scene.scene_number,
                (scene.heading or "")[:60],
                scene.int_ext or "–",
                scene.day_night or "–",
                str(scene.estimated_minutes or "–"),
            ])
        t = Table(table_data, colWidths=[50, 280, 50, 50, 70])
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1a1a2e")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
            ("FONTSIZE", (0, 0), (-1, -1), 9),
        ]))
        elements.append(t)
        elements.append(Spacer(1, 12))

    doc.build(elements)
    return buf.getvalue()


# ── Excel ─────────────────────────────────────────────────────────────────────

async def export_schedule_xlsx(project_id: str, db: AsyncSession) -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment

    project, days = await _load_schedule(project_id, db)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Shooting Schedule"

    header_fill = PatternFill("solid", fgColor="1a1a2e")
    header_font = Font(bold=True, color="FFFFFF")
    day_fill = PatternFill("solid", fgColor="e8e8f0")

    row = 1
    for day in days:
        ws.cell(row=row, column=1, value=f"Day – {day.date.strftime('%A, %d %b %Y')}")
        ws.cell(row=row, column=1).fill = day_fill
        ws.cell(row=row, column=1).font = Font(bold=True)
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=5)
        row += 1

        headers = ["Scene #", "Heading", "I/E", "D/N", "Est. (min)"]
        for col, h in enumerate(headers, 1):
            cell = ws.cell(row=row, column=col, value=h)
            cell.fill = header_fill
            cell.font = header_font
        row += 1

        for scene in day.scenes:
            ws.cell(row=row, column=1, value=scene.scene_number)
            ws.cell(row=row, column=2, value=scene.heading or "")
            ws.cell(row=row, column=3, value=scene.int_ext or "")
            ws.cell(row=row, column=4, value=scene.day_night or "")
            ws.cell(row=row, column=5, value=scene.estimated_minutes)
            row += 1
        row += 1

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── iCalendar ─────────────────────────────────────────────────────────────────

async def export_schedule_ical(project_id: str, db: AsyncSession) -> str:
    from icalendar import Calendar, Event

    project, days = await _load_schedule(project_id, db)
    cal = Calendar()
    cal.add("prodid", "-//SceneMind AI//EN")
    cal.add("version", "2.0")
    cal.add("calscale", "GREGORIAN")

    for day in days:
        event = Event()
        scenes_str = ", ".join(s.scene_number for s in day.scenes)
        event.add("summary", f"Shooting Day – Scenes {scenes_str}")
        event.add("dtstart", day.call_time or day.date)
        event.add("dtend", day.wrap_time_planned or day.date)
        event.add("description", f"Project: {project.name if project else project_id}")
        cal.add_component(event)

    return cal.to_ical().decode()
